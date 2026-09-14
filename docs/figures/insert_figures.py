# -*- coding: utf-8 -*-
"""Inserta las 5 figuras generadas en Tesis-Final 4.0.docx, con el mismo
formato de leyenda usado para las tablas (N en negrita, título en cursiva,
Nota en cursiva). NOTA: este script documenta la inserción original; ya fue
ejecutado una vez. Los numeros de figura (7, 9, 8 para las tres ultimas) y
los nombres de archivo reflejan la numeracion final del documento, que
quedo corrida respecto de la numeracion interna original de los scripts
gen_fig_*.py porque se insertaron 4 capturas de pantalla (Figuras 3-6)
entre la Figura 2 y lo que originalmente era "Figura 3" (hallazgo menor,
ronda 10 de auditoria independiente)."""
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

PATH = r'C:\Users\Acel\Desktop\Tesis Final\Tesis-Final 4.0.docx'
FIGDIR = r'C:\Users\Acel\Desktop\Tesis Final\TesisCiberseguridad\docs\figures'

d = docx.Document(PATH)


def find_para(needle):
    matches = [p for p in d.paragraphs if needle in p.text]
    if len(matches) != 1:
        raise SystemExit(f"ERROR: '{needle}' -> {len(matches)} matches")
    return matches[0]


def insert_figure_before(anchor_para, num, img_file, title, nota, width_in=5.8):
    # imagen (párrafo propio, centrado)
    img_p = anchor_para.insert_paragraph_before('')
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = img_p.add_run()
    run.add_picture(f'{FIGDIR}\\{img_file}', width=Inches(width_in))

    # "Figura N" en negrita (mismo estilo que "Tabla N")
    num_p = anchor_para.insert_paragraph_before('')
    num_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = num_p.add_run(f'Figura {num}')
    r.bold = True

    # título en cursiva
    title_p = anchor_para.insert_paragraph_before('')
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title_p.add_run(title)
    r.italic = True

    # Nota en cursiva
    nota_p = anchor_para.insert_paragraph_before('')
    nota_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = nota_p.add_run(f'Nota. {nota}')
    r.italic = True

    # separador en blanco antes del contenido siguiente
    anchor_para.insert_paragraph_before('')


# 1. Arquitectura de despliegue -> antes del heading 4.2
anchor1 = find_para('4.2. Construcción del Contenedor n8n Personalizado')
insert_figure_before(
    anchor1, 1, 'fig1_arquitectura.png',
    'Arquitectura de despliegue de dos capas del sistema',
    'Elaboración propia a partir de docker-compose.yml (repositorio del proyecto).'
)

# 2. Secuencia GMP -> antes del heading 4.5
anchor2 = find_para('4.5. El Pipeline de Análisis Completo en n8n')
insert_figure_before(
    anchor2, 2, 'fig2_secuencia_gmp.png',
    'Diagrama de secuencia del protocolo GMP entre el nodo Code y gvmd',
    'Elaboración propia a partir de docs/run_experiment.js (repositorio del proyecto).'
)

# 3. Matriz de confusion (heatmap) -> antes de "Tabla 3"
anchor3 = find_para('Tabla 3')
insert_figure_before(
    anchor3, 7, 'fig7_matriz_confusion.png',
    'Matriz de confusión de la clasificación de criticidad (mapa de calor)',
    'Elaboración propia a partir de docs/experiment_results.jsonl y docs/analyze_results.js (repositorio del proyecto).'
)

# 4. Flujo del algoritmo classifyAsset -> antes del inicio del codigo en Anexo E
anchor4 = find_para('Tabla de pesos por servicio para clasificacion')
insert_figure_before(
    anchor4, 9, 'fig9_flujo_algoritmo.png',
    'Diagrama de flujo del algoritmo de clasificación de criticidad (classifyAsset)',
    'Elaboración propia a partir de docs/classify_asset.js, verificado con docs/test_classify.js (repositorio del proyecto).'
)

# 5. Tiempos por repeticion -> antes de "Tabla 5"
anchor5 = find_para('Tabla 5')
insert_figure_before(
    anchor5, 8, 'fig8_tiempos_por_repeticion.png',
    'Tiempo total del ciclo por repetición, manual vs. automatizado',
    'Elaboración propia a partir de docs/experiment_results.jsonl (repositorio del proyecto).'
)

d.save(PATH)
print('OK: 5 figuras insertadas y documento guardado')
